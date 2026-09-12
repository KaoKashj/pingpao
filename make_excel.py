#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Convert the TSV results into .xlsx (multiple sheets) for Excel."""
import csv

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = "/Users/kaorouchuan/Documents/ChatGPT/选课"


def read_tsv(path):
    rows = []
    with open(path, encoding="utf-8") as f:
        r = csv.reader(f, delimiter="\t")
        for row in r:
            rows.append(row)
    return rows


def fill_sheet(ws, rows, widths=None):
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    for ci, cell in enumerate(rows[0], 1):
        c = ws.cell(row=1, column=ci, value=cell)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(vertical="center")
    for ri, row in enumerate(rows[1:], 2):
        for ci, val in enumerate(row, 1):
            ws.cell(row=ri, column=ci, value=val)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if widths:
        for ci, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(ci)].width = w


wb = Workbook()

# Sheet 1: 周五课程明细
rows1 = read_tsv(f"{BASE}/周五开课_候选课程明细.tsv")
ws1 = wb.active
ws1.title = "周五开课明细"
fill_sheet(ws1, rows1, widths=[13, 24, 16, 14, 26, 30, 18, 30, 22, 10, 26, 14, 8])

# Sheet 2: PDF 培养方案课程清单
rows2 = read_tsv(f"{BASE}/培养方案课程清单.tsv")
ws2 = wb.create_sheet("PDF培养方案55门")
fill_sheet(ws2, rows2, widths=[14, 30, 24])

out = f"{BASE}/周五课程结果.xlsx"
wb.save(out)
print("saved", out)
